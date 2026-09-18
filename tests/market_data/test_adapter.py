from dataclasses import replace
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from volcal.market_data.preprocessing import DataLoader
from volcal.market_data.schema import SPX_MID_SCHEMA, GOOGL_MID_SCHEMA, MarketDataError
from volcal.market_data.adapter import prepare_quotes
from volcal.utils import black


ROOT = Path(__file__).resolve().parents[2]


def ingest(asset, book, schema):
    return prepare_quotes(DataLoader(ROOT/'data'/asset, book).load_surface(schema, 'Mid'))


def test_spx_mid():
    result = ingest('spx', 'SPX_17_10_25.xlsx', SPX_MID_SCHEMA)
    q, df = result.quotes, result.observations
    assert not result.rejected
    assert len(df) == 77 and df.expiry.nunique() == 7
    assert result.source.valuation_date == pd.Timestamp('2025-10-17')
    assert q.S0 == (6543.93,)*77
    assert set(df.day_count) == {'ACT/365F'}
    first = df.iloc[0]
    assert first.original_header == '80%'
    assert first.r == .04005 and first.q == pytest.approx(.0057)
    assert first.market_iv == .3812
    assert first.F == 6581.58 and first.moneyness == .8
    assert first.K == .8*6581.58
    assert first['T'] == 61/365
    np.testing.assert_array_equal(df.K, df.moneyness*df.F)
    np.testing.assert_array_equal(q.F, df.F)
    assert set(df.moneyness) == {.8, .85, .9, .95, .975, 1., 1.025, 1.05, 1.1, 1.15, 1.2}
    assert set(df.loc[df.moneyness == 1., 'option_type']) == {'call'}
    assert np.all(df.loc[df.moneyness < 1., 'option_type'] == 'put')
    eighteen = df[df.expiry == '18M'].iloc[0]
    assert eighteen.F == 6829.11
    assert eighteen.forward_residual == pytest.approx(7.745303897656413)
    assert eighteen.F != eighteen.rates_forward
    assert df.forward_absolute_residual.max() == pytest.approx(7.745303897656413)


def test_googl_mid_exact_stored_strikes_and_units():
    result = ingest('googl', 'GOOGL_16_12_25.xlsx', GOOGL_MID_SCHEMA)
    q, df = result.quotes, result.observations
    assert not result.rejected
    assert len(df) == 153 and df.expiry.nunique() == 9
    assert result.source.valuation_date == pd.Timestamp('2025-12-16')
    assert q.S0 == (307.62,)*153
    assert set(df.day_count) == {'ACT/360'}
    first = df.iloc[0]
    assert first.r == .0413 and first.q == 0
    assert first.market_iv == .5423128762478799
    assert first.F == 308.43276135902113
    assert first.K == first.original_header == 156.86
    assert first['T'] == 23/360
    assert df['T'].max() == 1818/360 == 5.05
    np.testing.assert_array_equal(df.moneyness, df.K/df.F)
    np.testing.assert_array_equal(q.F, df.F)
    # Read raw cells independently: both headers and decimal data survive exactly.
    book = load_workbook(ROOT/'data/googl/GOOGL_16_12_25.xlsx', data_only=True)
    try:
        rows = list(book['Mid'].values)
    finally:
        book.close()
    headers = rows[0][7:]
    for raw in rows[1:10]:
        batch = df[df.expiry == raw[2]]
        np.testing.assert_array_equal(batch.K, headers)
        np.testing.assert_array_equal(batch.original_header, headers)
        np.testing.assert_array_equal(batch.market_iv, raw[7:])
        np.testing.assert_array_equal(batch.r, [raw[4]]*17)
        np.testing.assert_array_equal(batch.F, [raw[6]]*17)
    assert first.K != .5*first.F
    assert first.K != .5*first.S0
    # The 313.72 header stays a fixed strike across maturities, not ATM-forward.
    center = df[df.original_header == 313.72]
    assert np.all(center.K == 313.72)
    assert center.moneyness.nunique() == 9
    assert center.option_type.iloc[0] == 'call' and center.option_type.iloc[-1] == 'put'
    assert df.forward_absolute_residual.max() < 1e-12


@pytest.mark.parametrize('asset,book,schema', [
    ('spx', 'SPX_17_10_25.xlsx', SPX_MID_SCHEMA),
    ('googl', 'GOOGL_16_12_25.xlsx', GOOGL_MID_SCHEMA),
])
def test_complete_metadata_and_forward_black_prices(asset, book, schema):
    result = ingest(asset, book, schema)
    df = result.observations
    required = {'workbook', 'sheet', 'valuation_date', 'expiry', 'expiration_date',
        'original_header', 'T', 'day_count', 'S0', 'r', 'q', 'F', 'K', 'moneyness',
        'market_iv', 'option_type', 'market_price', 'market_vega', 'excel_row',
        'source_values', 'forward_residual', 'forward_relative_residual'}
    assert required <= set(df.columns)
    assert set(df.workbook) == {str(ROOT/'data'/asset/book)}
    assert set(df.sheet) == {'Mid'}
    np.testing.assert_allclose(df.market_price,
        black.price(df.market_iv, df['T'], df.K, F=df.F, r=df.r, option_type=df.option_type),
        atol=0, rtol=0)
    np.testing.assert_allclose(df.market_vega,
        black.vega(df.market_iv, df['T'], df.K, F=df.F, r=df.r), atol=0, rtol=0)
    # Round trip one near-ATM quote per maturity; prices include inconsistent carry.
    for _, batch in df.groupby('expiry'):
        row = batch.loc[(batch.moneyness-1).abs().idxmin()]
        recovered = black.iv_solver(row.market_price, row['T'], row.K,
            F=row.F, r=row.r, option_type=row.option_type)
        assert recovered == pytest.approx(row.market_iv, abs=1e-10)


def workbook(tmp_path, rows, headers=(100., 110.)):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Mid'
    sheet.append(['Act Date', 'Spot', 'Expiry', 'Exp Date', 'Risk Free', 'Impl (Yld)', 'ImplFwd', *headers])
    for row in rows:
        sheet.append(row)
    path = tmp_path/'quotes.xlsx'
    book.save(path)
    book.close()
    return DataLoader(tmp_path, path.name)


def test_rejections_are_auditable_and_forward_is_not_repaired(tmp_path):
    loader = workbook(tmp_path, [
        [pd.Timestamp('2025-01-01'), 70., '6M', pd.Timestamp('2025-07-01'), .03, .2, 115., .25, None],
        [None, None, '1Y', pd.Timestamp('2026-01-01'), .03, .2, -1., .25, .25],
        [None, None, 'expired', pd.Timestamp('2024-01-01'), .03, .2, 115., .25, .25],
    ])
    result = prepare_quotes(loader.load_surface(GOOGL_MID_SCHEMA))
    assert len(result.quotes.T) == 1
    assert len(result.rejected) == 5
    assert result.quotes.F == (115.,) and result.quotes.S0 == (70.,) and result.quotes.q == (.2,)
    assert result.observations.iloc[0].forward_absolute_residual > 40
    assert all(r.workbook.endswith('quotes.xlsx') and r.sheet == 'Mid' and r.reason for r in result.rejected)
    assert {r.excel_row for r in result.rejected} == {2, 3, 4}
    missing = next(r for r in result.rejected if r.excel_row == 2)
    assert missing.original_header == 110. and dict(missing.source_values)[110.] is None
    all_bad = loader.load_surface(GOOGL_MID_SCHEMA)
    all_bad.observations.loc[:, 'expiration_date'] = pd.Timestamp('2024-01-01')
    with pytest.raises(MarketDataError) as exc:
        prepare_quotes(all_bad)
    assert len(exc.value.rejected) == 6


def test_explicit_units_day_count_and_arbitrary_strike_header(tmp_path):
    loader = workbook(tmp_path, [
        [pd.Timestamp('2025-01-01'), 70., '6M', pd.Timestamp('2025-07-01'), .03, .2, 115., .25, .3],
    ], headers=(99.123456789, 112.987654321))
    first = prepare_quotes(loader.load_surface(GOOGL_MID_SCHEMA))
    changed = replace(GOOGL_MID_SCHEMA, day_count='ACT/365F', iv_unit='percentage_points')
    second = prepare_quotes(loader.load_surface(changed))
    assert first.quotes.K == second.quotes.K == (99.123456789, 112.987654321)
    assert first.quotes.market_iv == (.25, .3)
    assert second.quotes.market_iv == (.0025, .003)
    assert first.quotes.T == (181/360,)*2
    assert second.quotes.T == (181/365,)*2
    assert set(second.observations.day_count) == {'ACT/365F'}
    assert first.quotes.F == second.quotes.F == (115.,)*2


@pytest.mark.parametrize('field,value', [('quote_axis', 'spot_moneyness'),
    ('iv_unit', 'guess'), ('r_unit', 'guess'), ('q_unit', 'guess'), ('day_count', 'ACT/252')])
def test_invalid_schema(field, value):
    with pytest.raises(ValueError):
        replace(GOOGL_MID_SCHEMA, **{field: value})


@pytest.mark.parametrize('spot,date', [(71., '2025-01-01'), (70., '2025-01-02')])
def test_metadata_uniqueness_checked_independently(tmp_path, spot, date):
    loader = workbook(tmp_path, [
        [pd.Timestamp('2025-01-01'), 70., '6M', pd.Timestamp('2025-07-01'), .03, .2, 115., .25, .3],
        [pd.Timestamp(date), spot, '1Y', pd.Timestamp('2026-01-01'), .03, .2, 115., .25, .3],
    ])
    with pytest.raises(MarketDataError, match='exactly one'):
        loader.load_surface(GOOGL_MID_SCHEMA)


def test_existing_spx_loader_contract_unchanged():
    loader = DataLoader(ROOT/'data/spx', 'SPX_17_10_25.xlsx')
    spot, valuation, frame = loader.load_iv_table()
    assert spot == 6543.93 and valuation == pd.Timestamp('2025-10-17')
    assert list(frame.columns) == ['Expiry', 'Exp Date', 'Risk Free', 'ImplFwd', 'Impl (Yld)', 'Moneyness', 'IV']
    assert len(frame) == 77 and frame.iloc[0]['IV'] == .3812
    assert frame.iloc[0]['ImplFwd'] == 6581.58
